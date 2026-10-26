from pathlib import Path
import csv,gzip,itertools,json,textwrap,zipfile
import pandas as pd
ROOT=Path(__file__).resolve().parents[1]
PROJECT=ROOT/'gold_price_forecasting'
OUT=ROOT/'presentation'; OUT.mkdir(exist_ok=True)
TEAM='Priyanshu Modak (I080) | Vishesh Dokania (I019) | Parth Zanwar (I074)'
FX=json.loads((PROJECT/'data/usd_inr.json').read_text())
GRID=pd.read_csv(PROJECT/'outputs/arima_grid_search.csv')
SCORES=pd.read_csv(PROJECT/'outputs/model_comparison.csv',index_col=0)
SLIDES=[]
def slide(title,bullets,notes,image=None,tag='MAIN PRESENTATION'):
 SLIDES.append(dict(title=title,bullets=bullets,notes=notes,image=image,tag=tag))
slide('Gold Price Forecasting',[
 'Applied Time Series Analysis | B.Tech AI Semester VII | 2026-27',
 'ARIMA / SARIMA analysis with an INR-per-gram application',
 '18K, 20K, 22K and 24K fine-gold estimates',TEAM],
 'Good morning, Professor. Our project studies daily gold futures with the complete ATSA workflow. We first build and evaluate the model in USD per troy ounce, then convert forecasts into rupees per gram for four purities. We will show measured results and limitations. The analysis snapshot ends October 26, 2026.')
slide('The questions we investigated',[
 'Can past prices support useful daily forecasts?',
 'Does ARIMA or weekly SARIMA improve on tomorrow = today?',
 'How does uncertainty change with the forecast horizon?',
 'How can forecasts be displayed for Indian users?'],
 'This is a univariate time-series study with a demonstrable application. Success means correct preprocessing, a defensible evaluation and clear uncertainty. It does not require the most complex model to win. The Indian-price output estimates metal value, not a city-specific retail quote.')
slide('All 11 required ATSA steps',[
 '1 Collection | 2 EDA | 3 Missing values',
 '4 ADF/stationarity | 5 Transformations | 6 Seasonality',
 '7 ACF/PACF | 8 Model family | 9 Lags',
 '10 AIC/BIC | 11 Streamlit deployment'],
 'Map the work explicitly to the supplied slide. EDA includes line plots, scatter plots, histograms, QQ plots and descriptive statistics. Residual diagnostics, baseline evaluation and optional GARCH extend the checklist. The walkthrough identifies exact notebook sections and output files for each step.')
slide('Real data and provenance',[
 'Yahoo Finance GC=F: rolling COMEX gold futures',
 '2,971 source rows | 02 Jan 2015 to 26 Oct 2026',
 'Native price: USD per troy ounce',
 'USD-INR source: INR=X; dated snapshot retained'],
 'Show data/gold_daily.csv and PROVENANCE.md. The actual analysis uses real historical market observations. It is not Indian retail or MCX data. A rolling futures series can contain contract-roll effects, and current-session data may be provisional. FX is a separate dataset used for display conversion.')
slide('Architecture: analysis and application',[
 'Notebook: collect -> clean -> analyse -> select -> evaluate',
 'Artifacts: charts, CSVs, fitted model and metadata',
 'App: fetch/cache -> refit ARIMA -> forecast -> convert',
 'GitHub holds source; Streamlit provides hosting'],
 'The app reads the default order from metadata and refits the selected model on its selected history. It does not use the saved pickle as its main forecasting path. The notebook provides a separate saved-model helper. SARIMA comparison and GARCH are notebook analyses; the interactive app uses ARIMA. GitHub publication is verified. The user confirmed Streamlit deployment on October 27, 2026; the exact production URL was not supplied for independent verification.')
slide('Cleaning and chronological split',[
 'Sort/deduplicate; require positive numeric closes',
 'Forward-fill 111 missing weekdays: 3.60%',
 'Regular series: 3,082 business weekdays',
 'Train 2,982 through 08 Jun | Test 100 from 09 Jun 2026'],
 'Forward fill carries the last known price; no future price is used. Generic Monday-Friday indexing includes exchange holidays and introduces some zero returns. The hold-out consists of regularised weekdays, not strictly 100 observed trading sessions. Random splitting would mix past and future. The grid is fitted on training only.')
slide('EDA: price levels and dependence',[
 'Source mean: USD 1,966.46 per troy ounce',
 'Source-price standard deviation: USD 956.73',
 'Lagged prices move together; returns need separate study',
 'A rising price level alone does not establish predictability'],
 'Explain the level and rolling mean from the saved chart. These descriptive statistics use the 2,971 source observations before weekday filling. Consecutive price levels can be highly correlated simply because both share a stochastic trend. That is not proof of an exploitable forecasting signal. Lag scatter plots are saved separately.', '01_line_plot.png')
slide('EDA: distribution and heavy tails',[
 'Source-return excess kurtosis: 7.7843',
 'Jarque-Bera rejects normality; QQ tails depart',
 'Source-return annualised volatility: about 16.8%',
 'Heavy tails limit unqualified Gaussian interval claims'],
 'The QQ plot compares empirical return quantiles to normal quantiles. Tail departures suggest more extreme moves than the normal reference. Volatility is descriptive historical information, not a guaranteed future estimate. Normality failure matters especially for uncertainty assumptions; it does not automatically invalidate every point forecast.', '03_hist_qq.png')
slide('Stationarity: opposite null hypotheses',[
 'Price: ADF p=.9755; KPSS lower reported bound .01',
 'Log price: ADF p=.9897; KPSS lower bound .01',
 'Log return: ADF p about 6.83e-27; KPSS upper bound .10',
 'The sample supports first-differenced logs: d=1'],
 'ADF has a unit-root null; KPSS has a stationarity null. Small ADF and non-small KPSS support stationarity in this sample. ADF non-rejection does not prove a unit root. A printed zero is rounded, not a literally zero probability. KPSS reports tabulated bounds. Mean-process stationarity can coexist with changing conditional variance.')
slide('Transformation and selected equation',[
 'z(t) = log(P(t)); r(t) = z(t) - z(t-1)',
 'ARIMA(p,1,q) models short-memory dynamics in r(t)',
 'Selected: z(t) = z(t-1) + drift + innovation',
 'exp(log forecast) is a conditional median'],
 'Logs express proportional changes and keep inverse forecasts positive. First differences remove the integrated level. With p=q=0 this is a log random walk with drift. A mean on the price scale would require a distribution-appropriate bias correction, such as exp(mu+v/2) for Gaussian log uncertainty. The app reports the exponentiated location without that correction.', '04_transformations.png')
slide('Seasonality: examine rather than assume',[
 'Strength: 5 days .0002 | 21 days .0084',
 '252-business-day strength: .1747',
 'Annual decomposition is descriptive evidence only',
 'A weekly SARIMA competitor is evaluated'],
 'A 252-day period approximates an annual business-day cycle rather than an exact calendar year. Classical decomposition imposes a periodic structure and uses centred calculations, so it is exploratory here. We did not evaluate every seasonal order or annual SARIMA. The actual forecasting competitor uses seasonal order (1,0,1,5).', '05_decomposition.png')
slide('ACF/PACF and model identification',[
 'AR: PACF cutoff; ACF usually decays',
 'MA: ACF cutoff; PACF usually decays',
 'ARMA: both decay; seasonal spikes motivate SARIMA',
 'Squared-return dependence motivates variance analysis'],
 'These are finite-sample heuristics, not automatic proofs. Some spikes occur by chance across many inspected lags. AR and MA terms act on the differenced log process inside ARIMA. Moving-average innovation terms are not the same as a rolling mean. The squared-return plot concerns volatility rather than just mean predictability.', '07_acf_pacf.png')
slide('The recorded 25-model search',[
 'p=0..4; d=1; q=0..4 -> 25 candidate orders',
 'Same training log-price series; drift in every candidate',
 'AIC selects (1,1,4); BIC selects (0,1,0)',
 'Convergence warnings occurred: rankings need a caveat'],
 'AIC is 2k minus twice log likelihood; BIC is k log n minus twice log likelihood. BIC penalises complexity more at this sample size. Negative scores are valid. The notebook keeps finite scores and does not persist a convergence flag for every candidate. Therefore do not claim 25 certified converged fits or an optimum over all possible ARIMA models.', '08_aic_bic_heatmap.png')
slide('Two evaluation designs',[
 'One-step: forecast, then observe for later predictions',
 'Model parameters stay fixed during hold-out updates',
 'Fixed-origin: forecast 100 dates with no test updates',
 'The app refits available history and forecasts 5..60 weekdays'],
 'Appending test observations with refit=False and obtaining one-step predictions updates model states with past observations while retaining training parameters. This differs from forecasting the whole test path at once. The reported RMSE comparison is one-step; do not call it 30-day-ahead RMSE. The separate fixed-origin chart addresses the harder multi-step problem.', '09_test_forecasts.png')
slide('Measured performance: naive wins RMSE',[
 'Naive 60.5337 | BIC ARIMA 60.6723 USD/oz',
 'AIC ARIMA 60.6684 | SARIMA 60.6085 USD/oz',
 'BIC MAPE 1.0778%; directional agreement 49%',
 'HAC loss-difference p=.4639: BIC ARIMA vs naive'],
 'State the result directly. Naive has the lowest RMSE, while AIC ARIMA has slightly lower MAE/MAPE. Different metrics can rank models differently. BIC is a prespecified simple-model choice, not the hold-out RMSE winner. A non-small p-value fails to detect a loss difference; it does not prove equality or market efficiency. MAPE is not classification accuracy.')
slide('Intervals and residual diagnostics',[
 'Fixed-origin 100-day interval coverage: 94/100',
 'Residual Ljung-Box p=.08397 at lag 10; .000589 at lag 20',
 'ARCH-LM p about 6.28e-45; normality rejected',
 'Longer-lag dependence and changing variance remain'],
 'This is a single correlated path, not independent proof of interval calibration. Residual diagnostics use the training fit with the first initialisation residual removed. It would be false to say all assumptions passed. Longer-lag correlation, squared-residual dependence and non-normality limit strong adequacy claims. The nominal 95% bands depend on model assumptions.', '10_multistep_forecast.png')
slide('Optional GARCH volatility extension',[
 'Variance = omega + alpha*past shock^2 + beta*past variance',
 'Training alpha=.0359; beta=.9583; sum=.9942',
 'Student-t innovations fitted to percentage log returns',
 'Final normal-quantile bands remain approximate'],
 'ARIMA models conditional mean; GARCH models changing conditional variance. Persistence near one indicates slow decay in the fitted variance process. The final band sums predicted daily variances and uses 1.96 around the ARIMA log forecast. It is not an exact Student-t interval or a jointly simulated ARIMA-GARCH system. The interactive app does not offer a GARCH engine.', '12_garch_volatility.png')
slide('Final saved forecast',[
 'Selected ARIMA refitted on all 3,082 observations',
 '30-business-day snapshot saved as a forecast CSV',
 '27 Oct 2026 median: USD 4,039.29 per troy ounce',
 'ARIMA interval: USD 3,957.84 to 4,122.41'],
 'This is the saved October 26 analysis snapshot. A later live app can differ because of new data or the default ten-year app history versus the notebook start in 2015. It is not a contradiction. Demonstrate outputs/forecast_next_30_days.csv and explain the difference between a point median and an uncertainty interval.', '13_final_forecast.png')
slide('INR per gram and all four karats',[
 'INR/g = USD/oz x INR/USD / 31.1034768 x fineness',
 '18K .750 | 20K .833 | 22K .916 | 24K .999',
 'Dated automatic FX or a positive manual override',
 'Same factor scales history, median and both bounds'],
 'Precious metals use a troy ounce, not the ordinary ounce. We use linked BIS fineness values: 22K=.916 rather than exactly 22/24 and 24K fine gold=.999 rather than exactly one. These are four purity scalings of one underlying forecast, not four independently trained models. Units and conversion assumptions are visible in the app.')
slide('What the Indian-price display means',[
 'Selected FX is held fixed for all future dates',
 'History is revalued at that same selected rate',
 'Future FX risk, GST, duties and making charges excluded',
 'Metal-value estimates are not jewellery-shop quotations'],
 'The displayed history is not an observed historical INR series. We would need date-aligned historical FX and future FX modelling to build that. At fixed USD gold prices a 10% rise in INR per USD scales every INR value by 10%, without changing native-model AIC/BIC. Location, premiums, taxes and making charges explain differences from a retail quote.')
slide('Professor demo: deliberate actions',[
 'Start 24K, 30 weekdays, ARIMA(0,1,0)',
 'Switch 22K -> 20K -> 18K; inspect comparison table',
 'Compare horizons 5 and 60; change FX and order',
 'Export CSVs; explain valid upload and real-data fallback'],
 'Predict each effect before clicking. Purity scales the output but does not invent a new price pattern. Longer horizons create additional rows and generally broader uncertainty for the selected random walk. Changing history or order triggers a model change and can be slow. Show saved results at the end, including the naive baseline and failed diagnostic tests. Do not delete data to stage an outage.')
slide('All finite setting combinations',[
 '4 karats x 13 history values x 56 horizons x 72 orders',
 '72 orders = 6 p choices x 2 d choices x 6 q choices',
 '209,664 discrete tuples are fully enumerated',
 'Enumeration is not 209,664 executed model fits'],
 'History ranges 3..15, horizon 5..60, p/q 0..5 and d 1 or 2. The compressed CSV lists all finite setting tuples. Continuous FX inputs, arbitrary upload contents, future observations and network conditions cannot be exhaustively enumerated. The scenario matrix selects useful boundary and operational cases. High-order combinations may fail or be impractical.')
slide('Decision combinations and evidence',[
 'ADF rejects + KPSS does not: stationarity supported',
 'ADF does not reject + KPSS rejects: investigate differencing',
 'Both reject or neither rejects: evidence is inconclusive',
 'ARCH rejection: inspect variance, not just mean dynamics'],
 'Avoid automatic test-based buttons. Examine trend specification, sample size, structural breaks and diagnostics. If AIC improves but hold-out RMSE worsens, complexity may not generalise. Repeatedly tuning to the same hold-out biases conclusions. The walkthrough includes decision tables, the complete order catalogue and professor Q&A.')
slide('Code walkthrough and deployment',[
 'streamlit_app.py: root cloud entry point',
 'app.py: loaders -> ARIMA -> INR conversion -> exports',
 'run_analysis.py and make_report.py: analysis/report generation',
 'Public GitHub source; Streamlit deployment confirmed by user'],
 'The root requirements file is the runtime stack; the inner file also includes notebook dependencies. File paths are anchored to app.py. The app is deployed on Streamlit Community Cloud according to the user. Temporary deployment buttons were removed after deployment. Use the actual production URL in the viva; no URL is invented here. The walkthrough maps functions to their actual behaviour.')
slide('Limitations and next steps',[
 'No demonstrated RMSE advantage over naive',
 'Convergence flags and rolling-origin checks need improvement',
 'No forecast of FX, interest rates, inflation or news',
 'Next: aligned INR target, trading calendar, calibrated intervals'],
 'Separate completed work from future work. Proposed improvements include storing and filtering convergence outcomes, multiple rolling origins, horizon-specific accuracy, aligned FX and an exchange-aware calendar. Prophet, LSTM, annual SARIMA and exogenous regressors are not implemented and have no results in this project.')
slide('Conclusion and handover',[
 'Complete ATSA workflow on preserved real market data',
 'd=1 supported; BIC favours a drifted random walk',
 'Measured errors and diagnostic failures reported honestly',
 'Working INR/g app with four purity views'],
 'Our contribution is a reproducible statistical workflow and an understandable application. We can demonstrate provenance, preprocessing, model selection, baseline evaluation and currency/purity arithmetic. Every team member should explain the code and actual contributions, and disclose AI assistance in line with faculty requirements.')
slide('Backup: order families and trend rules',[
 '(0,1,0): random walk with log drift',
 '(p,1,0): AR in differences | (0,1,q): MA in differences',
 '(p,1,q): ARMA in first-differenced log prices',
 'App d=2 is available; its deterministic trend is n'],
 'For d=1 the app uses trend t, giving drift in first differences. For d=2 it uses no deterministic trend. So d comparisons change both differencing and trend assumptions. d=0, seasonal order controls and GARCH selection are not offered in the app. The 72-order table records these rules.',tag='BACKUP / VIVA')
slide('Backup: failure and fallback combinations',[
 'Gold live fails + real cache: labelled cached prices',
 'FX live fails + snapshot: dated saved exchange rate',
 'No FX source: require a positive manual rate',
 'Bad CSV / insufficient data: stop rather than fabricate'],
 'Evidence labels distinguish verified UI checks from code-reviewed or planned scenarios. Required CSV columns are Date and Close; units must be USD per troy ounce. Invalid numeric and nonpositive prices are filtered; at least 100 valid observations are needed. Dates are sorted and duplicates keep the last. Fit warnings are not always hard failures.',tag='BACKUP / VIVA')
slide('Backup: four questions to answer precisely',[
 'Does 1.1% MAPE mean 98.9% accuracy?',
 'Why can naive outperform fitted ARIMA?',
 'Why does the app differ from the notebook?',
 'Why is a 24K result different from a jeweller quote?'],
 'Answers: MAPE is average relative error, not classification accuracy. Weak predictable structure and estimation error can make naive competitive. A new window, order, data date or FX setting changes the app output. Jewellery quotes include local/tax/product effects absent from a futures-derived metal estimate. Never turn these results into unsupported trading claims.',tag='BACKUP / VIVA')
slide('References and assistance disclosure',[
 'Yahoo Finance GC=F and INR=X; source dates disclosed',
 'statsmodels, arch and Forecasting: Principles and Practice',
 'NIST conversion and BIS fineness grades',
 'AI-assisted implementation; state actual team contributions'],
 'Primary references are linked in the walkthrough. The course PDF states a GenAI allowance and citation expectations. The implementation and this pack used AI assistance. Do not claim unaided authorship or automatic compliance with a numerical policy limit. This is an educational project.',tag='REFERENCES')
SECTIONS=[]
def section(title,paragraphs=None,rows=None,headers=None):
 SECTIONS.append(dict(title=title,paragraphs=paragraphs or [],rows=rows,headers=headers))
section('1. How to use the professor pack',[
 'Main narrative: slides 1-26, approximately 18-22 minutes, plus a five-minute demo. Slides 27-29 are backups; slide 30 gives references. For ten minutes use slides 1,2,3,4,6,9,13,15,18,19,21,25,26.',
 'The PPTX includes speaker notes. SPEAKER_SCRIPT.md provides the full spoken script. The walkthrough explains code, results, combinations, demonstration steps and viva answers. Results are from the saved October 26, 2026 run; the live app can differ.',
 'All 209,664 finite setting combinations are enumerated, not batch-tested. Evidence labels distinguish measured results, reviewed code paths, and planned demonstrations.'])
section('2. An opening you can speak',[
 'Good morning, Professor. Our ATSA mini project analyses daily gold futures through stationarity testing, decomposition, autocorrelation and ARIMA/SARIMA modelling. We use a chronological hold-out and compare against a naive benchmark. We then provide rupees-per-gram estimates for 18K, 20K, 22K and 24K gold. A simple BIC-selected random walk with drift is competitive but does not improve the naive baseline in RMSE. We report the uncertainty and residual limitations honestly.'])
section('3. Team and presentation handoffs',[
 TEAM,
 'Suggested speaking allocation: Priyanshu introduces the problem, provenance and preprocessing; Vishesh explains statistical modelling and evaluation; Parth demonstrates the app and limitations. These are suggested speaking roles, not claims about who wrote the code. Replace with actual contributions.',
 'Implementation and these materials used AI assistance. Explain every component and disclose actual contributions; do not assume this pack automatically meets the course PDF\'s stated GenAI policy.'])
section('4. Mapping every required ATSA step',rows=[
 ['1','Collection','Notebook section 1; data/gold_daily.csv and PROVENANCE.md'],
 ['2','EDA','Notebook section 2; line/scatter/histogram/QQ; figures 01-03'],
 ['3','Missing values','Notebook section 3; 111 weekdays forward-filled'],
 ['4','Stationarity / ADF','Notebook section 4; stationarity.csv and KPSS'],
 ['5','Transformations','Notebook section 5; logs and difference; figure 04'],
 ['6','Seasonality','Notebook section 6; strengths; figures 05-06'],
 ['7','ACF / PACF','Notebook section 7; figure 07'],
 ['8','Family identification','Notebook section 8; AR/MA/ARMA/ARIMA/SARIMA'],
 ['9','Lags','Notebook section 9; p/q grid'],
 ['10','AIC / BIC','arima_grid_search.csv and figure 08'],
 ['11','Deployment','streamlit_app.py, inner app.py, GitHub, cloud form']],headers=['Step','Requirement','Evidence'])
section('5. Dataset: what it is and what it is not',[
 'Source is Yahoo Finance GC=F, a rolling COMEX gold futures series. The snapshot has 2,971 observations from January 2, 2015 to October 26, 2026, with OHLCV fields. Close is the target and its native unit is USD per troy ounce.',
 'It is not Indian city-specific retail, MCX or jewellery-shop data. Contract-roll effects can affect a futures series and a current-session value may be provisional. The actual analysis uses real observations; simulated test fixtures were used only in separate app checks.',
 f'FX source is INR=X. The saved rate is {FX["rate"]:.6f} INR per USD dated {FX["date"]}. It is a dated observation, not a permanently valid current rate.',
 'Raw-source mean is 1,966.45857 USD/oz, standard deviation 956.72956, minimum 1,050.80005 and maximum 5,318.39990. Source log-return excess kurtosis is 7.7843 and annualised volatility about 16.8%. These statistics precede weekday filling.'])
section('6. Preprocessing, missing dates and leakage',[
 'Notebook processing sorts dates, retains the last duplicate, converts Close to numeric, replaces invalid or nonpositive prices, reindexes to B frequency, forward-fills and drops leading missing values. It adds 111 missing weekdays (3.60%), yielding 3,082 observations.',
 'Forward fill uses past information. Backward fill and two-sided interpolation can inject future observations. Generic weekdays include exchange holidays and introduce zero returns; this is an approximation to disclose, not an exact trading calendar.',
 'App uploads require Date and Close, parseable dates, positive numeric prices and at least 100 valid source observations. The app sorts and deduplicates. Uploaded values must be fine-gold USD per troy ounce; semantic units cannot be inferred from a numeric column.',
 'Training is 2,982 regular observations through June 8, 2026; the hold-out is 100 regular weekdays from June 9 to October 26. Candidate fitting uses training only. Full-sample exploratory plots must not become a reason to repeatedly tune against future observations.'])
section('7. Explain every EDA plot',rows=[
 ['Line plot','Changing level and rolling mean','A trend is not proof of usable predictive information'],
 ['Lag scatter','Price(t) versus Price(t-1), then returns','Integrated levels can look highly dependent'],
 ['Histogram','Distribution, skewness and tails','Price and return distributions differ'],
 ['QQ plot','Empirical versus normal quantiles','Tail departures limit Gaussian assumptions'],
 ['Descriptive statistics','Counts, moments and volatility','State whether raw or regularised observations are used']],headers=['Output','Explanation','Caution'])
section('8. Stationarity and all ADF/KPSS outcomes',[
 'Weak stationarity concerns a stable mean and variance with autocovariance depending on lag. Mean-process stationarity does not imply constant conditional variance at every date.',
 'ADF null is a unit root; KPSS null is level stationarity (regression c here). ADF uses AIC-based lag selection; KPSS uses automatic lag selection. The 5% threshold is an aid, not proof.'],rows=[
 ['Reject','Do not reject','Stationarity supported for the tested series'],
 ['Do not reject','Reject','Non-stationarity evidence; investigate transformation/differencing'],
 ['Reject','Reject','Conflicting; inspect trend, breaks and test specifications'],
 ['Do not reject','Do not reject','Inconclusive; limited power may matter']],headers=['ADF','KPSS','Interpretation'])
section('9. Actual stationarity evidence',[
 'Price: ADF p=.9755, KPSS p at lower reported bound .01. Log price: ADF p=.9897, KPSS lower bound .01. Log return: ADF about 6.83e-27, KPSS upper reported bound .10. The saved CSV rounds a very small ADF p to zero; it is not literally a zero probability.',
 'These results support first-differenced logs for the mean process in this sample. ADF failure to reject does not prove a unit root. KPSS tabulated bounds are not exact tail probabilities. A second-difference stationarity result would not prove first differences were insufficient.'])
section('10. Transformation and ARIMA mathematics',[
 'z_t = log(P_t); r_t = z_t - z_(t-1). Logs express proportional changes and preserve positive inverse forecasts. First differencing addresses the integrated level; it does not guarantee Gaussian or homoskedastic errors.',
 'ARIMA(p,d,q) differences its target d times and models the result using p autoregressive and q moving-average innovation terms. The target supplied here is log price. SARIMA adds seasonal (P,D,Q,s).',
 'Selected ARIMA(0,1,0) with drift: z_t = z_(t-1) + c + e_t. Its h-step log location is z_T + h*c. The app reports exp(mu_h), a conditional median under the log model. A Gaussian conditional price mean would require exp(mu_h + v_h/2), which is not the default.',
 'The app sets trend t for d=1, producing drift after differencing. For d=2 it uses trend n, so a d=2 comparison changes deterministic terms as well. App d=0 is not available.'])
section('11. Seasonality and decomposition',[
 'Classical additive decomposition of log price examines periods 5, 21 and 252 business days. Recorded strengths are .0002, .0084 and .1747. Weekly/monthly strengths are weak; annual descriptive strength is more noticeable.',
 'A 252-business-day annual approximation is not exact calendar seasonality. Decomposition imposes a periodic structure and uses centred operations. It is exploratory here, not a future-aware predictor. The monthly-return ANOVA is another descriptive check.',
 'Only the weekly SARIMA competitor with seasonal order (1,0,1,5) is forecast and evaluated. Annual SARIMA and all seasonal permutations were not searched.'])
section('12. ACF/PACF: all model-family cases',rows=[
 ['AR(p)','PACF cutoff; ACF decay','AR terms model past values of a stationary/differenced process'],
 ['MA(q)','ACF cutoff; PACF decay','Past innovations; not a rolling average'],
 ['ARMA','Both generally decay','AR and MA in stationary differences'],
 ['ARIMA','Differencing plus ARMA','Our main family; logs supplied as target'],
 ['SARIMA','Additional seasonal dynamics','Weekly competitor only'],
 ['GARCH','Squared-return dependence / ARCH','Conditional variance rather than mean trend']],headers=['Family','Heuristic','Meaning'])
section('13. Model selection and all 25 actual candidates',[
 'The notebook tries all p and q in {0,1,2,3,4}, d fixed at 1: exactly 25 tuples, each with drift and the same training target. ACTUAL_25_ARIMA_RESULTS.csv contains the recorded scores, sorted by BIC.',
 'Minimum AIC is ARIMA(1,1,4), approximately -18852.29. Minimum BIC is ARIMA(0,1,0), approximately -18836.86. AIC=2k-2 log L and BIC=k log n-2 log L. Compare scores only on comparable likelihood/data settings. Negative scores are valid.',
 'Convergence warnings occurred in the executed notebook. The grid retains finite scores without persisting/filtering every convergence flag. Therefore the recorded ranking is not a claim that every candidate converged or a global optimum over all possible orders.'])
section('14. Evaluation designs and information availability',[
 'One-step evaluation fits training parameters and updates the state using previous hold-out observations with refit=False. Each prediction conditions on the past, not its own actual target. Appending observations and asking for one-step predictions is distinct from forecasting the whole future block.',
 'Fixed-origin evaluation forecasts all 100 dates from the training endpoint without incorporating hold-out actuals. The separate chart shows that harder problem. Reported comparison RMSE is one-step, not a 30-day fixed-origin performance estimate.',
 'After evaluation the selected order is refitted on the full series for final outputs. Exploratory analysis uses the full sample; repeatedly changing choices after looking at plots or test scores would bias the hold-out. A prespecified grid and multiple rolling origins would strengthen the study.'])
section('15. Metrics and recorded results',[
 'RMSE=sqrt(mean(error^2)); MAE=mean(abs(error)); MAPE=100*mean(abs(error)/actual). RMSE emphasises large misses. Gold prices are positive, so the usual zero-denominator MAPE issue is limited here. MAPE is not classification accuracy.',
 'Directional accuracy compares the signs of actual and predicted changes. The code excludes zero predicted directions; naive predicts no change so its directional score is undefined.'],rows=[
 [name,f'{r.RMSE:.4f}',f'{r.MAE:.4f}',f'{r["MAPE %"]:.4f}',f'{r["Directional acc %"]:.1f}' if pd.notna(r['Directional acc %']) else 'Undefined'] for name,r in SCORES.iterrows()],headers=['Model','RMSE USD/oz','MAE USD/oz','MAPE %','Direction %'])
section('16. Significance and conclusions to defend',[
 'Naive has the lowest RMSE (60.5337 versus BIC ARIMA 60.6723), while AIC ARIMA has slightly lower MAE and MAPE than the other methods. Metric rankings can differ. BIC selection is prespecified simplicity, not a test-RMSE win.',
 'The loss-difference comparison regresses squared-error differences on a constant with a HAC standard error. BIC ARIMA versus naive p=.46395; AIC ARIMA p=.73392; SARIMA p=.70141. It is asymptotic and does not prove equality, market efficiency or universal forecasting quality.',
 'Do not say 1.0778% MAPE equals 98.9222% accuracy. Do not say ARIMA significantly beats naive. If operational selection were solely the observed hold-out RMSE, naive would be preferred in this comparison.'])
section('17. Residual tests and all mean/variance cases',[
 'Training-fit residuals drop the first diffuse-initialisation residual. Ljung-Box tests serial dependence; ARCH-LM tests time-varying variance; Jarque-Bera tests normality.',
 'Actual p-values: Ljung-Box residual lag 10 .083968; lag 20 .0005889; squared residual lag 10 4.40e-71; ARCH-LM 6.28e-45; normality effectively underflows to zero. It is incorrect to say the residuals are fully white noise or all assumptions pass.'],rows=[
 ['No residual correlation','No ARCH evidence','Mean/variance checks appear adequate; validate coverage too'],
 ['Residual correlation','No ARCH evidence','Revisit mean dynamics/lags/omissions'],
 ['No residual correlation','ARCH evidence','Investigate conditional variance'],
 ['Residual correlation','ARCH evidence','Both mean and variance need investigation']],headers=['Mean diagnostic','Variance diagnostic','Response'])
section('18. Intervals and optional GARCH',[
 'Fixed-origin nominal 95% ARIMA bands contain 94/100 hold-out observations. This is a single correlated path, not proof of calibrated repeated coverage at each horizon. Marginal intervals are not simultaneous guarantees for the whole future path.',
 'Training GARCH(1,1), fitted to percentage log returns with Student-t errors, has rounded omega .0092385, alpha .0359, beta .9583 and persistence .9942. Persistence near one indicates slowly decaying fitted variance shocks.',
 'The final GARCH band uses predicted daily variances, cumulative standard deviation and 1.96 around an ARIMA log forecast. It is approximate, not exact Student-t inference or a jointly estimated ARIMA-GARCH path simulation. Its coverage is not established. The current app shows ARIMA intervals only.'])
section('19. Rupees, grams and every karat',[
 'INR/g = USD/oz x INR per USD / 31.1034768 x fineness. Use a troy ounce rather than the ordinary ounce. Implemented fineness values are 18K .750, 20K .833, 22K .916 and 24K fine gold .999.',
 'All four purity views scale one underlying gold model. The same positive factor must scale the median, history and both bounds. 22K=.916 is not exactly 22/24; 24K=.999 is the fine-gold convention explicitly used in the app.',
 'Illustration only: the table assumes USD 4,000/oz and FX 100 INR/USD; these are teaching inputs, not current quotations.'],rows=[[k,str(f),f'{4000*100/31.1034768*f:,.2f}'] for k,f in {'18K':.750,'20K':.833,'22K':.916,'24K':.999}.items()],headers=['Karat','Fineness','Illustrative INR/g'])
section('20. Every FX path and price interpretation',[
 f'Live INR=X is tried first; a valid rate and date are displayed. If it fails, the saved JSON supplies {FX["rate"]:.6f} INR/USD dated {FX["date"]}. If both fail, the app requires a positive manual rate. A manual override replaces conversion only; it does not forecast FX.',
 'Selected FX is held constant across all historical and future displayed dates. History therefore revalues old USD prices at the selected rate; it is not a historically aligned INR series. A 10% FX increase scales every INR value by 10% while leaving native ARIMA fitting and AIC/BIC unchanged.',
 'Displayed metal values exclude GST, duties, making charges, local premiums and city-specific effects. Future exchange-rate uncertainty is not included in the bands. A jewellery-shop quote is not a directly comparable evaluation target.'])
section('21. Every app control',rows=[
 ['Purity','18K/20K/22K/24K','Scales display, not independent training'],
 ['History','3..15 years','Changes sample; live/calendar and cached/252 approximations differ'],
 ['Horizon','5..60 weekdays','Changes count of forecast dates, not calendar days'],
 ['p','0..5','AR order after differencing'],
 ['d','1 or 2','Difference order; trend t vs n'],
 ['q','0..5','MA innovation order'],
 ['FX','Positive automatic/manual','Changes display-unit conversion'],
 ['CSV','Date/Close in USD per troy ounce','At least 100 valid prices; years slider does not trim uploads']],headers=['Control','Range','Effect'])
section('22. All permutations: what can be exhaustive',[
 'Finite combinations: 4 purities x 13 history lengths x 56 horizons x (6 p-values x 2 d-values x 6 q-values) = 209,664. ALL_209664_APP_COMBINATIONS.csv.gz lists each tuple once, labelled ENUMERATED_NOT_BATCH_TESTED.',
 'ALL_72_MODEL_ORDERS.csv lists every app order, its trend rule and notebook-grid inclusion. ACTUAL_25_ARIMA_RESULTS.csv is the measured grid. Enumeration does not prove convergence, successful fitting, accuracy or deployment of each combination.',
 'Arbitrary positive FX values, arbitrary CSV contents, future data, calendar edge cases and network conditions have no useful finite exhaustive catalogue. The 26-case scenario matrix selects meaningful paths and labels evidence.'])
section('23. A five-minute professor demo',[
 '1. Run .\\.venv\\Scripts\\python.exe -m streamlit run streamlit_app.py. Open http://localhost:8501. State the gold source/date and FX source/date.',
 '2. Begin 24K, 10 years, 30 weekdays, (0,1,0). Explain last close, median, interval, AIC/BIC and the futures-derived estimate qualification.',
 '3. Switch to 22K, 20K and 18K. Predict the price ordering first; show the comparison table. Explain that returns/pattern shape do not change under fixed positive scaling.',
 '4. Compare horizons 5 and 60. Check row count and uncertainty. Generally increasing width is expected for the selected random walk; do not promise strict monotonic width for every arbitrary order.',
 '5. Set manual FX to 100; demonstrate proportional change. Restore automatic FX. Native-model fit scores are not rupee-denominated accuracy scores.',
 '6. Try (1,1,0), explain AR in returns, then optionally (0,1,1). Return to (0,1,0). Avoid slow high-order stress cases in the main demo.',
 '7. Download selected-purity and all-purity CSVs. Show explicit units/karat/FX. If uploading, prepare a valid USD/oz file; do not upload INR/g data.',
 '8. Open model_comparison.csv and the residual diagnostics. Close with the naive RMSE advantage and remaining assumptions. Do not hide these results.',
 '9. Show the public GitHub source and open your actual deployed Streamlit URL. Deployment was confirmed by the user on October 27, 2026; temporary deployment buttons have been removed.'])
section('24. Edge cases and evidence discipline',[
 'DEMO_AND_EDGE_CASE_MATRIX.csv contains 26 paths. Previously verified runtime checks include all four purities, manual FX arithmetic, ordered bounds, a mocked real-data fallback and an earlier d=2 smoke test. Other rows are labelled code review or demo plan rather than executed checks.',
 'Bad required columns or dates show an error. Invalid numeric/nonpositive values are removed; insufficient remaining prices stop the fit. High-order fits may warn, fail or run slowly. Warnings do not always prevent an output.',
 'Do not delete real files to create a viva outage. Cache corruption and empty metadata are not fully hardened production paths. A numeric INR/g upload cannot be automatically recognised as wrong units.'])
section('25. File-by-file and function-by-function walkthrough',rows=[
 ['streamlit_app.py','runpy executes inner app','Root cloud entry point'],
 ['app.py constants','Troy grams and FINENESS','Defines conversion semantics'],
 ['load_exchange_rate','Live INR=X, then JSON','Returns positive FX, date and source'],
 ['load_prices','Live GC=F, then real CSV','No automatic fabricated prices'],
 ['fit_forecast','log/B-frequency/ffill/ARIMA','Caches median and bounds'],
 ['sidebar and upload','Purity/FX/history/horizon/order','User sensitivity and validated input'],
 ['conversion and chart','Positive factor on all series','Current-FX revaluation'],
 ['exports and comparison','One or all purities','CSV units, karat and FX metadata'],
 ['notebook sections 1-9','Collection to selection','EDA/tests and training grid'],
 ['notebook sections 10-11','Evaluation and diagnostics','One-step, fixed-origin, GARCH'],
 ['notebook section 12','Refit/save/helper','Pickle helper is separate from app refit'],
 ['run_analysis.py','NotebookClient, then report','Creates executed notebook'],
 ['make_report.py','Reads CSVs/metadata','Generates RESULTS.md'],
 ['requirements files','Root runtime, inner analysis','Pinned runtime; notebook stack separate']],headers=['File/block','Operation','Explain'])
section('26. Architecture questions to answer',[
 'Why not use the saved pickle in every app request? The app permits new history and order choices, so it refits; metadata chooses defaults. The notebook includes a separate persistence helper.',
 'Does changing order run the grid? No, it fits the selected order. The archived notebook RMSE appears only for matching default ticker/order and is labelled archived. Changing history can still make the new fit different.',
 'Why can live and cached histories differ? Live requests use calendar-year periods; cached slices use years*252 observations. Uploads use supplied history without applying the years slider. These are current implementation limitations.',
 'What hosting is verified? GitHub is public and local health was checked. The user confirmed Streamlit deployment on October 27, 2026. The exact production URL has not been supplied or independently checked.'])
section('27. Prioritised improvements',[
 'First record convergence flags and exclude failed candidates. Next prespecify rolling-origin tests over multiple horizons, periods and data regimes. Then align historical FX and consider an INR target or joint FX scenarios.',
 'Use an exchange-aware calendar, evaluate horizon-specific interval coverage and consider heavy-tail simulations. Add exogenous factors only with proper time-aligned information availability.',
 'Improve upload/history consistency, archived-metric comparability and versioned data/model provenance. Annual SARIMA, Prophet, LSTM and external drivers are future work, not implemented results.'])
QA=[
 ('Why choose gold?','It is a real positive financial series with changing levels, dependence and volatility, suitable for the full ATSA workflow.'),
 ('Is the data real?','Yes. The GC=F snapshot has 2,971 historical Yahoo observations. Simulated data was used only in separate UI checks, not the reported analysis.'),
 ('Why 3,082 rows after cleaning?','Reindexing creates 111 missing weekdays, filled from past prices. These are not 111 extra real trading observations.'),
 ('Why use Close?','The defined target is the daily closing price. OHLCV fields are available but are not external regressors.'),
 ('Why chronological splitting?','A random split can train on future information relative to a target; chronology matches forecasting information availability.'),
 ('Why log prices?','Logs express proportional changes and give positive inverse forecasts. They do not guarantee stationarity or normality.'),
 ('What is d?','The number of differences applied to log prices. The selected model uses one.'),
 ('What is p?','The autoregressive lag order in the differenced log process.'),
 ('What is q?','The number of lagged innovation terms. It is not a rolling-average window.'),
 ('Why ADF plus KPSS?','Opposite null hypotheses provide a useful cross-check, although neither is definitive.'),
 ('Does ADF p>.05 prove a unit root?','No; it fails to reject under that specification and sample, with power and structural-break limitations.'),
 ('Does p=0 mean certainty?','No. It is rounded/underflowed. The saved log-return ADF summary is about 6.83e-27.'),
 ('Can stationary returns have volatility clustering?','Yes. Mean-process stationarity testing and time-varying conditional variance address different questions.'),
 ('Why p and q only through four in the notebook?','It is a tractable bounded design motivated by short-lag dynamics, not proof that no higher order can help.'),
 ('Did all candidates converge?','Not established. Warnings occurred and per-order convergence flags were not persisted. Scores need that caveat.'),
 ('Why does BIC favour simplicity?','Its complexity penalty is k log n; at this sample size small likelihood gains may not offset additional parameters.'),
 ('Can AIC/BIC be negative?','Yes. Compare relative values from comparable likelihoods and datasets.'),
 ('Why not repeatedly tune on test RMSE?','That makes the test set part of model selection and biases its evaluation. Use a validation/rolling-origin design.'),
 ('What is ARIMA(0,1,0) with drift?','A random walk plus constant expected log change, with innovations; not a deterministic trading rule.'),
 ('Why test SARIMA when seasonality is weak?','It supplies a concrete weekly competitor rather than assuming seasonal terms must help.'),
 ('Was annual SARIMA tested?','No. Annual decomposition was descriptive; the evaluated seasonal competitor is (1,0,1,5).'),
 ('Did ARIMA beat naive?','Not by RMSE. Naive is 60.5337 versus BIC ARIMA 60.6723 USD/oz on the recorded one-step hold-out.'),
 ('Is MAPE equivalent to accuracy?','No. It is average relative regression error, not the probability of a correct classification.'),
 ('Is direction prediction reliable?','The BIC model has 49% directional agreement in this hold-out, so useful directional advantage is not demonstrated.'),
 ('What does the HAC comparison show?','At p about .464 for BIC ARIMA, it does not detect a squared-loss difference. It does not prove equality or efficiency.'),
 ('Does a 95% band cover the whole future path with 95% probability?','No. Marginal forecast intervals and simultaneous path coverage are different.'),
 ('Was interval coverage measured?','94 of 100 fixed-origin actual values lie within ARIMA bands. One correlated path does not establish general calibration.'),
 ('Why GARCH?','Squared-residual dependence and ARCH rejection indicate changing conditional variance. GARCH models variance, not the price trend.'),
 ('Are the final GARCH bands exact Student-t intervals?','No. Student-t innovations were fitted but the final calculation uses normal quantiles and cumulative variance.'),
 ('Did all residual assumptions pass?','No. Lag-20 Ljung-Box, ARCH and normality reject, limiting adequacy claims.'),
 ('Why INR per gram?','It is familiar to Indian users; conversion is separate from the native USD/oz statistical fitting.'),
 ('Why divide by 31.1034768?','That is grams per troy ounce. Precious metals use troy rather than ordinary ounces.'),
 ('Why .916 rather than exactly 22/24?','The app explicitly adopts the BIS-linked 916 fineness grade.'),
 ('Is 24K exactly pure in this implementation?','No. It is fine gold 999, or 99.9%, as stated in the interface.'),
 ('Are there four independently trained purity models?','No. A single underlying forecast is scaled by fineness.'),
 ('Is future USD-INR predicted?','No. A selected dated or manual rate is fixed over forecast dates and FX uncertainty is excluded.'),
 ('Why does the value differ from the jeweller?','It is a converted futures metal value; tax, local premiums, duties and making charges are excluded.'),
 ('Why does the app differ from the notebook?','History length, order, new observations and FX may differ. Notebook metrics are archived and labelled.'),
 ('What happens offline?','Default paths use the real cached gold CSV and dated FX snapshot. Missing sources lead to required manual FX or an error, not fabricated observations.'),
 ('How many combinations were evaluated?','The notebook records 25 nonseasonal candidate scores and the stated forecast competitors. The 209,664 app tuples are enumerated, not benchmarked.'),
 ('What improvement comes first?','Persist convergence flags, strengthen rolling-origin evaluation, then align FX/calendar handling and validate intervals.'),
 ('Who wrote the project?','State actual team contributions and AI assistance. Do not invent individual coding roles or claim numerical policy compliance without checking.')]
section('28. Viva: 42 questions and precise answers',[f'{i}. {q}\nAnswer: {a}' for i,(q,a) in enumerate(QA,1)])
section('29. Claims to avoid',[
 'Avoid 99% accuracy. Say: BIC ARIMA MAPE was 1.0778% on the specified one-step hold-out.',
 'Avoid ARIMA beats naive. Say: naive wins RMSE; AIC ARIMA has slightly lower MAE/MAPE.',
 'Avoid all assumptions pass. Say: longer-lag dependence, ARCH effects and non-normality remain.',
 'Avoid four trained karat models. Say: one gold forecast, four fineness scalings.',
 'Avoid exact Indian retail price. Say: futures-derived INR/g metal-value estimate excluding retail charges.',
 'Avoid every combination tested. Say: finite settings enumerated; actual score and check evidence labelled.',
 'Avoid inventing a production URL. Say: the user confirmed Streamlit deployment; show the actual live address from your account.'])
section('30. Reproducibility and useful commands',[
 'Local app: .\\.venv\\Scripts\\python.exe -m streamlit run streamlit_app.py',
 'Analysis: .\\.venv\\Scripts\\python.exe gold_price_forecasting\\run_analysis.py',
 'Rebuild pack: .\\.venv\\Scripts\\python.exe scripts\\build_professor_pack.py',
 'Pack-builder libraries: python-pptx, python-docx, reportlab plus the project dependencies. Preserve the dated CSV, metadata, charts and executed notebook for an unchanged viva; downloading again may change results.'])
section('31. Primary references and disclosure',[
 'Yahoo gold data: https://finance.yahoo.com/quote/GC=F/history/ ; FX: https://finance.yahoo.com/quote/INR=X/history/',
 'statsmodels ARIMA: https://www.statsmodels.org/stable/generated/statsmodels.tsa.arima.model.ARIMA.html ; time-series API: https://www.statsmodels.org/stable/tsa.html',
 'Hyndman and Athanasopoulos, Forecasting: Principles and Practice (third edition): https://otexts.com/fpp3/',
 'arch forecasting: https://arch.readthedocs.io/en/stable/univariate/forecasting.html',
 'BIS fineness listing: https://lims.bis.gov.in/home/search_is_number/?is_number__doc_no=1417',
 'NIST conversion: https://www.nist.gov/pml/special-publication-811/nist-guide-si-appendix-b-conversion-factors/nist-guide-si-appendix-b8',
 'Streamlit: https://docs.streamlit.io/deploy/streamlit-community-cloud/deploy-your-app/deploy',
 'Implementation, documentation and this pack used AI assistance. Team members should understand and verify the work and discuss the course PDF\'s stated GenAI policy with faculty. This pack makes no unsupported compliance claim.'])

SCENARIOS=[
 ('D01','Default 24K, 10y, 30 days, order 0/1/0','Forecast and ordered bounds','Verified UI path'),
 ('D02','Switch every karat','Only fineness scale changes','Verified all four'),
 ('D03','Manual FX=100','USD/oz *100 /31.1034768 *fineness','Verified arithmetic'),
 ('D04','d=2','No deterministic trend fit; inspect shape','Earlier smoke test verified'),
 ('D05','Horizon 5 vs 60','Check row counts and uncertainty','Demo plan'),
 ('D06','History 3 vs 15','Different sample and refit','Demo plan'),
 ('D07','Order 1/1/0','AR structure in returns','Demo plan'),
 ('D08','Order 0/1/1','MA innovations in returns','Demo plan'),
 ('D09','Order 1/1/1','ARMA in returns','Demo plan'),
 ('D10','Order 5/2/5','May be slow/warn/fail','Stress plan; no success claim'),
 ('D11','Gold live unavailable, cache present','Label cached real data','Mocked outage verified'),
 ('D12','FX live unavailable, snapshot present','Label saved rate/date','Mocked outage verified'),
 ('D13','No FX sources','Require positive manual FX','Code review only'),
 ('D14','No gold sources','Stop; no fake data','Code review only'),
 ('D15','CSV missing Date/Close','Reject required columns','Code review only'),
 ('D16','Malformed CSV dates','Invalid CSV message','Code review only'),
 ('D17','Invalid/zero/negative prices','Filter; require 100 remaining','Code review only'),
 ('D18','Duplicate/unsorted dates','Sort; keep last duplicate','Code review only'),
 ('D19','Fewer than 100 valid prices','Reject insufficient sample','Code review only'),
 ('D20','CSV is INR/g not USD/oz','Semantic mismatch cannot be inferred','Documented user contract'),
 ('D21','Selected-purity CSV','Check units, karat and FX metadata','Code/UI review; inspect download in demo'),
 ('D22','All-purity CSV','Four forecast columns and FX','Code/UI review'),
 ('D23','Change order away from default','Hide archived RMSE','Code review'),
 ('D24','New live prices','Outputs may differ from snapshot','Expected behaviour'),
 ('D25','GitHub public access','Anonymous repo available','Verified remotely'),
 ('D26','Deployed app UI','Temporary deploy buttons absent; forecasts remain usable','User confirmed deployment; buttons removed')]
pd.DataFrame(SCENARIOS,columns=['id','scenario','expected','evidence']).to_csv(OUT/'DEMO_AND_EDGE_CASE_MATRIX.csv',index=False)
with gzip.open(OUT/'ALL_209664_APP_COMBINATIONS.csv.gz','wt',encoding='utf-8',newline='') as f:
 w=csv.writer(f); w.writerow(['id','karat','history_years','horizon_business_days','p','d','q','status'])
 for i,v in enumerate(itertools.product(['18K','20K','22K','24K'],range(3,16),range(5,61),range(6),[1,2],range(6)),1): w.writerow([i,*v,'ENUMERATED_NOT_BATCH_TESTED'])
assert i==209664
orders=[]
for p,d,q in itertools.product(range(6),[1,2],range(6)):
 family='Random walk' if p==q==0 and d==1 else 'AR in differences' if p>0 and q==0 else 'MA in differences' if p==0 and q>0 else 'ARMA in differences' if p>0 and q>0 else 'Twice-integrated noise'
 orders.append([p,d,q,'t' if d==1 else 'n',family,d==1 and p<=4 and q<=4])
pd.DataFrame(orders,columns=['p','d','q','app_trend','interpretation','in_notebook_grid']).to_csv(OUT/'ALL_72_MODEL_ORDERS.csv',index=False)
GRID.sort_values('BIC').to_csv(OUT/'ACTUAL_25_ARIMA_RESULTS.csv',index=False)
(OUT/'presentation_content.json').write_text(json.dumps(SLIDES,indent=2),encoding='utf-8')
md=['# ATSA Gold Price Forecasting: Full Professor Walkthrough','',TEAM,'','Evidence snapshot: October 26, 2026.','']
for s in SECTIONS:
 md+=['## '+s['title'],'']
 for p in s['paragraphs']: md+=[p,'']
 if s['rows']:
  md+=['| '+' | '.join(s['headers'])+' |','| '+' | '.join(['---']*len(s['headers']))+' |']
  md+=['| '+' | '.join(str(v).replace('|','/') for v in row)+' |' for row in s['rows']]+['']
(OUT/'FULL_PROJECT_WALKTHROUGH.md').write_text('\n'.join(md),encoding='utf-8')
script=['# Slide-by-slide speaker script','',TEAM,'']
for i,s in enumerate(SLIDES,1): script+=[f'## Slide {i}: {s["title"]}','',s['notes'],'']
(OUT/'SPEAKER_SCRIPT.md').write_text('\n'.join(script),encoding='utf-8')
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from xml.sax.saxutils import escape
from docx import Document
NAVY='14283F'; GOLD='B68A32'; INK='263448'; PALE='F6F3EB'; MUTED='65758A'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
pdf=canvas.Canvas(str(OUT/'ATSA_Gold_Presentation.pdf'),pagesize=(960,540))
def ppt_text(sl,x,y,w,h,text,size=20,color=INK,bold=False):
 box=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame
 tf.word_wrap=True; tf.margin_left=0; tf.margin_top=0
 for j,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph(); p.text=line
  p.font.name='Aptos'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color); p.space_after=Pt(8)
 return box

def pdf_text(c,text,x,y,width,font=19,color=INK,leading=26):
 c.setFillColor(HexColor('#'+color)); c.setFont('Helvetica',font)
 for line in textwrap.wrap(text,width=max(15,int(width/(font*.52)))):
  c.drawString(x,y,line); y-=leading
 return y

for num,item in enumerate(SLIDES,1):
 sl=prs.slides.add_slide(prs.slide_layouts[6]); sl.background.fill.solid(); sl.background.fill.fore_color.rgb=RGBColor.from_string(PALE)
 bar=sl.shapes.add_shape(1,0,0,prs.slide_width,Inches(.12)); bar.fill.solid(); bar.fill.fore_color.rgb=RGBColor.from_string(GOLD); bar.line.fill.background()
 ppt_text(sl,.5,.32,12.3,.35,item['tag'],11,GOLD,True)
 ppt_text(sl,.5,.85,12.3,.85,item['title'],30,NAVY,True)
 has_image=item['image'] is not None
 y=2.0
 for b in item['bullets']:
  ppt_text(sl,.65,y,5.0 if has_image else 11.9,1.03,'\u2022 '+b,18 if has_image else 22)
  y+=1.08
 if has_image:
  from PIL import Image
  path=PROJECT/'figures'/item['image']
  with Image.open(path) as im: iw,ih=im.size
  scale=min(6.85/iw,4.7/ih); nw,nh=iw*scale,ih*scale
  sl.shapes.add_picture(str(path),Inches(6.0+(6.85-nw)/2),Inches(1.9+(4.7-nh)/2),width=Inches(nw),height=Inches(nh))
 ppt_text(sl,.5,7.03,11,.25,'ATSA | Gold Price Forecasting | Saved analysis: 26 Oct 2026',10,MUTED)
 ppt_text(sl,12.2,7.03,.65,.25,f'{num:02d}',11,GOLD,True)
 sl.notes_slide.notes_text_frame.text=item['notes']
 pdf.setFillColor(HexColor('#'+PALE)); pdf.rect(0,0,960,540,fill=1,stroke=0)
 pdf.setFillColor(HexColor('#'+GOLD)); pdf.rect(0,531,960,9,fill=1,stroke=0)
 pdf.setFont('Helvetica-Bold',10); pdf.drawString(36,505,item['tag'])
 pdf_text(pdf,item['title'],36,465,885,font=26,color=NAVY,leading=31)
 y=390
 for b in item['bullets']: y=pdf_text(pdf,'- '+b,46,y,355 if has_image else 860,font=15 if has_image else 19,leading=23)-17
 if has_image: pdf.drawImage(ImageReader(str(PROJECT/'figures'/item['image'])),430,85,width=495,height=310,preserveAspectRatio=True,anchor='c',mask='auto')
 pdf.setFont('Helvetica',9); pdf.setFillColor(HexColor('#'+MUTED)); pdf.drawString(36,24,'ATSA | Gold Price Forecasting | Saved analysis: 26 Oct 2026'); pdf.drawRightString(923,24,f'{num:02d}'); pdf.showPage()
prs.save(OUT/'ATSA_Gold_Presentation.pptx'); pdf.save()

styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyGold',fontName='Helvetica',fontSize=10,leading=15,spaceAfter=9))
styles.add(ParagraphStyle(name='HeadGold',fontName='Helvetica-Bold',fontSize=15,leading=19,textColor=HexColor('#'+NAVY),spaceBefore=14,spaceAfter=10,keepWithNext=True))
styles.add(ParagraphStyle(name='CellGold',fontSize=8,leading=11,splitLongWords=True))
story=[Paragraph('Gold Price Forecasting',styles['Title']),Paragraph('Full professor walkthrough and viva preparation',styles['Heading2']),Paragraph(escape(TEAM),styles['BodyGold']),Paragraph('Applied Time Series Analysis | B.Tech AI Semester VII | 2026-27',styles['BodyGold']),Paragraph('Evidence snapshot: October 26, 2026. Methodology, real results, code, settings, demo cases and 42 viva answers.',styles['BodyGold']),PageBreak()]
doc=Document(); doc.add_heading('Gold Price Forecasting',0); doc.add_heading('Full Professor Walkthrough',1); doc.add_paragraph(TEAM); doc.add_paragraph('ATSA | B.Tech AI Semester VII | 2026-27 | Snapshot: 26 October 2026')
for s in SECTIONS:
 story.append(Paragraph(escape(s['title']),styles['HeadGold'])); doc.add_heading(s['title'],1)
 for p in s['paragraphs']:
  story.append(Paragraph(escape(p).replace('\n','<br/>'),styles['BodyGold'])); doc.add_paragraph(p)
 if s['rows']:
  count=len(s['headers']); widths=[(A4[0]-86)/count]*count
  if count==3: widths=[115,145,A4[0]-86-260]
  cells=[[Paragraph(escape(str(v)),styles['CellGold']) for v in row] for row in [s['headers']]+s['rows']]
  t=Table(cells,colWidths=widths,repeatRows=1,hAlign='LEFT')
  t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),HexColor('#E8DFC8')),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.35,HexColor('#D6DCE4')),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
  story+=[t,Spacer(1,10)]
  dt=doc.add_table(rows=1,cols=count); dt.style='Light Shading Accent 1'
  for j,h in enumerate(s['headers']): dt.rows[0].cells[j].text=h
  for row in s['rows']:
   dcells=dt.add_row().cells
   for j,v in enumerate(row): dcells[j].text=str(v)
def footer(c,d):
 c.setFont('Helvetica',8); c.setFillColor(HexColor('#'+MUTED)); c.drawString(43,25,'ATSA Gold Price Forecasting | Actual results: 26 Oct 2026'); c.drawRightString(A4[0]-43,25,str(d.page))
SimpleDocTemplate(str(OUT/'FULL_PROJECT_WALKTHROUGH.pdf'),pagesize=A4,leftMargin=43,rightMargin=43,topMargin=42,bottomMargin=43).build(story,onFirstPage=footer,onLaterPages=footer)
doc.save(OUT/'FULL_PROJECT_WALKTHROUGH.docx')
(OUT/'README.md').write_text('''# Professor presentation pack

- ATSA_Gold_Presentation.pptx: 30 slides with embedded speaker notes.
- ATSA_Gold_Presentation.pdf: portable slide version.
- FULL_PROJECT_WALKTHROUGH.pdf / .docx / .md: methodology, code, decisions, demo, and 42 viva answers.
- SPEAKER_SCRIPT.md: slide-by-slide spoken guidance.
- ACTUAL_25_ARIMA_RESULTS.csv: recorded notebook model scores.
- ALL_72_MODEL_ORDERS.csv: every app order and deterministic-trend rule.
- ALL_209664_APP_COMBINATIONS.csv.gz: exhaustive finite settings catalogue, not batch-tested fits.
- DEMO_AND_EDGE_CASE_MATRIX.csv: 26 operational cases with evidence status.
- presentation_content.json: slide content snapshot.

Main slides: 1-26; backups 27-29; references 30. Results are from October 26, 2026.
Live app values can change with data/history/order/FX. User confirmed Streamlit deployment on October 27, 2026; no production URL is invented.
Rebuild: python scripts/build_professor_pack.py (python-pptx, python-docx, reportlab required).
''',encoding='utf-8')
with zipfile.ZipFile(ROOT/'ATSA_Professor_Presentation_Pack.zip','w',zipfile.ZIP_DEFLATED) as z:
 for f in OUT.iterdir():
  if f.is_file(): z.write(f,'presentation/'+f.name)
 z.write(Path(__file__),'scripts/build_professor_pack.py')
print('Generated:',len(SLIDES),'slides,',len(SECTIONS),'sections,',len(QA),'viva answers, 209664 tuples,',len(SCENARIOS),'scenarios.')
